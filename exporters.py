import csv
from datetime import datetime
from pathlib import Path
from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import config
from scrapers.base import Product

class Exporter:
    @staticmethod
    def export_all(products: List[Product], base_name: str = "promocoes_pet") -> dict:
        """Exporta a lista de produtos para Excel, Bloco de Notas (TXT) e CSV."""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        excel_path = config.OUTPUT_DIR / f"{base_name}_{timestamp}.xlsx"
        txt_path = config.OUTPUT_DIR / f"{base_name}_{timestamp}.txt"
        csv_path = config.OUTPUT_DIR / f"{base_name}_{timestamp}.csv"

        Exporter.to_excel(products, excel_path)
        Exporter.to_notepad_txt(products, txt_path)
        Exporter.to_csv(products, csv_path)

        return {
            "excel": str(excel_path),
            "txt": str(txt_path),
            "csv": str(csv_path)
        }

    @staticmethod
    def to_notepad_txt(products: List[Product], file_path: Path):
        """Exporta para um formato de Bloco de Notas claro, direto e fácil de ler."""
        with open(file_path, "w", encoding="utf-8") as f:
            f.write("=" * 75 + "\n")
            f.write("         🐾 PROMOÇÕES DE BANHO E TOSA / PETSHOP ENCONTRADAS 🐾\n")
            f.write(f"  Gerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M:%S')}\n")
            f.write(f"  Total de ofertas filtradas: {len(products)}\n")
            f.write("=" * 75 + "\n\n")

            if not products:
                f.write("Nenhuma promoção atendendo aos critérios mínimos foi encontrada.\n")
                return

            for i, p in enumerate(products, 1):
                f.write(f"[{i}] {p.title}\n")
                f.write(f"    🏪 Loja: {p.store}\n")
                
                if p.original_price and p.original_price > p.current_price:
                    f.write(f"    💰 Preço: R$ {p.current_price:,.2f}  (De: R$ {p.original_price:,.2f})  📉 DESCONTO: -{p.discount_percent:.0f}%\n".replace(",", "X").replace(".", ",").replace("X", "."))
                else:
                    f.write(f"    💰 Preço: R$ {p.current_price:,.2f}  📉 DESCONTO: -{p.discount_percent:.0f}%\n".replace(",", "X").replace(".", ",").replace("X", "."))

                stars = "★" * int(round(p.rating)) + "☆" * (5 - int(round(p.rating)))
                f.write(f"    ⭐ Avaliação: {p.rating:.1f}/5.0 ({p.reviews_count} avaliações) {stars}\n")
                if p.free_shipping:
                    f.write("    🚚 Frete Grátis: Sim\n")
                f.write(f"    🔗 Link: {p.url}\n")
                f.write("-" * 75 + "\n")

    @staticmethod
    def to_excel(products: List[Product], file_path: Path):
        """Exporta para planilha Excel (.xlsx) com layout moderno e links clicáveis."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Promoções Petshop"

        headers = [
            "Loja",
            "Produto",
            "Preço Promocional (R$)",
            "Preço De (R$)",
            "Desconto (%)",
            "Avaliação",
            "Nº Avaliações",
            "Frete Grátis",
            "Link da Oferta"
        ]

        # Estilo do Cabeçalho
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border_thin = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Inserção das linhas de dados
        discount_highlight = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

        for row_idx, p in enumerate(products, start=2):
            row_data = [
                p.store,
                p.title,
                p.current_price,
                p.original_price if p.original_price else "",
                f"{p.discount_percent:.1f}%",
                p.rating,
                p.reviews_count,
                "Sim" if p.free_shipping else "Não",
                p.url
            ]
            ws.append(row_data)

            # Estilização das células
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = border_thin
                c.font = Font(name="Calibri", size=10)

                # Formatações numéricas
                if col_idx in (3, 4):  # Preços
                    c.number_format = 'R$ #,##0.00'
                    c.alignment = Alignment(horizontal="right")
                elif col_idx in (1, 5, 6, 7, 8):
                    c.alignment = Alignment(horizontal="center")

                # Destaque se tiver desconto >= 20%
                if col_idx == 5 and p.discount_percent >= 20.0:
                    c.fill = discount_highlight
                    c.font = Font(name="Calibri", size=10, bold=True, color="276A3C")

                # Link clicável
                if col_idx == 9 and p.url:
                    c.hyperlink = p.url
                    c.font = Font(name="Calibri", size=10, color="0563C1", underline="single")

        # Ajuste automático da largura das colunas
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if cell.column == 9:  # Link
                    val_str = "Acessar Oferta"
                    cell.value = "Abrir Link"
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

        wb.save(file_path)

    @staticmethod
    def to_csv(products: List[Product], file_path: Path):
        """Exporta para CSV com codificação utf-8-sig compatível com Excel."""
        with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow([
                "Loja", "Produto", "Preco_Atual", "Preco_Original",
                "Desconto_Pct", "Avaliacao", "Num_Avaliacoes", "Frete_Gratis", "Link"
            ])
            for p in products:
                writer.writerow([
                    p.store,
                    p.title,
                    f"{p.current_price:.2f}",
                    f"{p.original_price:.2f}" if p.original_price else "",
                    f"{p.discount_percent:.1f}",
                    f"{p.rating:.1f}",
                    p.reviews_count,
                    "Sim" if p.free_shipping else "Nao",
                    p.url
                ])
